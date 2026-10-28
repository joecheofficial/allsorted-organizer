#!/usr/bin/env python3
"""Bounded, read-only local document reader. One JSON result, no network."""
import json
import os
import pathlib
import shutil
import subprocess
import sys
import tempfile
import zipfile

MAX_INPUT = 64 * 1024 * 1024
MAX_EXPANDED = 80 * 1024 * 1024


def extract_document(file, limit, max_units, ocr):
    result = dict(status="extracted", text="", units=0, ocrUnits=0, truncated=False, error=None)
    parts = []
    length = 0

    def append(text):
        nonlocal length
        text = str(text or "").replace("\x00", "")
        remaining = limit - length
        if len(text) > remaining:
            result["truncated"] = True
        if remaining > 0:
            chunk = text[:remaining]
            parts.append(chunk)
            length += len(chunk) + 1

    if os.path.getsize(file) > MAX_INPUT:
        result.update(status="too_large", error="Document exceeds the 64 MiB input limit.")
        return result
    ext = pathlib.Path(file).suffix.lower()
    if ext in (".docx", ".pptx", ".xlsx"):
        with zipfile.ZipFile(file) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED:
                result.update(status="too_large", error="Office archive exceeds the expanded size limit.")
                return result
    if ext == ".pdf":
        from pypdf import PdfReader
        from pypdf.errors import DependencyError
        try:
            reader = PdfReader(file)
        except DependencyError:
            # PdfReader can inspect encryption before exposing is_encrypted.
            # An unavailable AES backend is not evidence of a damaged PDF.
            result.update(status="missing_dependency", error="Encrypted PDF inspection requires an optional AES crypto reader. Supply an unlocked copy from its owner; no password recovery is attempted.")
            return result
        if reader.is_encrypted:
            result.update(status="encrypted", error="Encrypted PDF requires the owner to supply an unlocked copy.")
            return result
        result["truncated"] = len(reader.pages) > max_units
        for index, page in enumerate(reader.pages[:max_units]):
            if length >= limit:
                result["truncated"] = True
                break
            text = page.extract_text() or ""
            if not text.strip() and ocr:
                if not shutil.which("pdftoppm") or not shutil.which("tesseract"):
                    result.update(status="missing_dependency", error="OCR needs pdftoppm and tesseract on PATH.")
                    break
                with tempfile.TemporaryDirectory(prefix="brain-ocr-") as temporary:
                    prefix = os.path.join(temporary, "page")
                    env = dict(os.environ, OMP_THREAD_LIMIT="1")
                    subprocess.run(["pdftoppm", "-f", str(index + 1), "-l", str(index + 1), "-singlefile", "-scale-to", "1600", "-png", file, prefix], check=True, capture_output=True, timeout=30, env=env)
                    read = subprocess.run(["tesseract", prefix + ".png", "stdout"], check=True, capture_output=True, timeout=30, env=env)
                    text = read.stdout.decode("utf-8", errors="replace")
                    result["ocrUnits"] += 1
            append("[Page " + str(index + 1) + "]\n" + text if text.strip() else "")
            result["units"] += 1
            if not text.strip():
                result["truncated"] = True
    elif ext == ".docx":
        from docx import Document
        doc = Document(file)
        for index, paragraph in enumerate(doc.paragraphs):
            if index >= max_units or length >= limit:
                result["truncated"] = True
                break
            append(paragraph.text)
            result["units"] += 1
        for table in doc.tables:
            for row in table.rows:
                if result["units"] >= max_units or length >= limit:
                    result["truncated"] = True
                    break
                append(" | ".join(cell.text for cell in row.cells))
                result["units"] += 1
    elif ext == ".pptx":
        from pptx import Presentation
        deck = Presentation(file)
        result["truncated"] = len(deck.slides) > max_units
        for index, slide in enumerate(deck.slides):
            if index >= max_units or length >= limit:
                result["truncated"] = True
                break
            append("[Slide " + str(index + 1) + "]")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    append(shape.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        append(" | ".join(cell.text for cell in row.cells))
            if slide.has_notes_slide:
                append(slide.notes_slide.notes_text_frame.text)
            result["units"] += 1
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(file, read_only=True, data_only=False, keep_links=False)
        rows = 0
        try:
            for sheet in book.worksheets:
                if length >= limit or rows >= max_units:
                    result["truncated"] = True
                    break
                append("[Sheet " + sheet.title + "]")
                result["units"] += 1
                for row in sheet.iter_rows(values_only=True):
                    if rows >= max_units or length >= limit:
                        result["truncated"] = True
                        break
                    append(" | ".join(str(cell) if cell is not None else "" for cell in row))
                    rows += 1
        finally:
            book.close()
    else:
        raise ValueError("Unsupported document format")
    result["text"] = "\n".join(parts).strip()
    if not result["text"] and result["status"] == "extracted":
        result.update(status="empty", error="No document text was recovered; empty pages are not covered.")
    return result


def main():
    try:
        # Bound CPU and open descriptors. The parent also enforces wall time.
        import resource
        resource.setrlimit(resource.RLIMIT_CPU, (90, 90))
        resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
        if sys.platform.startswith("linux"):
            resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
    except (ImportError, ValueError, OSError):
        pass
    try:
        file, limit, units, ocr = sys.argv[1:]
        result = extract_document(file, min(max(int(limit), 1), 2_000_000), min(max(int(units), 1), 10000), ocr == "true")
    except ModuleNotFoundError as error:
        result = dict(status="missing_dependency", text="", units=0, ocrUnits=0, truncated=False, error="Missing local Python reader: " + str(error.name))
    except (subprocess.TimeoutExpired, subprocess.CalledProcessError):
        result = dict(status="failed", text="", units=0, ocrUnits=0, truncated=False, error="Local OCR failed or exceeded its time limit.")
    except Exception as error:
        result = dict(status="corrupt", text="", units=0, ocrUnits=0, truncated=False, error="Document could not be parsed (" + type(error).__name__ + ").")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
